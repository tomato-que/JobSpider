"""简历文件文本提取与敏感信息打码。

支持文本型 PDF / DOCX / TXT / MD；不做 OCR，扫描版 PDF 会明确报错。
原文件不落盘，调用方只保存提取后的文本。
"""
from __future__ import annotations

import io
import re
from pathlib import PurePath

MAX_UPLOAD_BYTES = 8 * 1024 * 1024
MAX_RESUME_CHARS = 20000
MAX_PDF_PAGES = 30
ALLOWED_EXTS = {".pdf", ".docx", ".txt", ".md"}


class ResumeParseError(ValueError):
    """简历无法解析时抛出，message 可直接展示给用户。"""


_ID_RE = re.compile(r"(?<!\d)\d{17}[\dXx](?!\d)")
_PHONE_RE = re.compile(r"(?<!\d)1[3-9]\d{9}(?!\d)")
_EMAIL_RE = re.compile(
    r"[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}")


def normalize_text(text: str) -> str:
    """统一换行、清理行尾空白、压缩连续空行，并限制总长度。"""
    text = (text or "").replace("\r\n", "\n").replace("\r", "\n")
    lines = [line.rstrip() for line in text.split("\n")]
    text = "\n".join(lines)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if len(text) > MAX_RESUME_CHARS:
        text = text[:MAX_RESUME_CHARS] + "\n...(已截断)"
    return text


def mask_contacts(text: str) -> str:
    """隐藏手机号、邮箱、18 位身份证号；简历解析不需要这些信息。"""
    if not text:
        return ""
    text = _ID_RE.sub("[证件号已隐藏]", text)
    text = _PHONE_RE.sub("[手机号已隐藏]", text)
    text = _EMAIL_RE.sub("[邮箱已隐藏]", text)
    return text


def _extract_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as e:
        raise ResumeParseError(
            "缺少 pypdf 依赖，请执行 pip install -r requirements.txt") from e

    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception as e:
        raise ResumeParseError(f"PDF 读取失败：{type(e).__name__}") from e

    if getattr(reader, "is_encrypted", False):
        try:
            reader.decrypt("")
        except Exception:
            pass
        if getattr(reader, "is_encrypted", False):
            raise ResumeParseError("PDF 已加密，无法提取文本")

    parts: list[str] = []
    for idx, page in enumerate(reader.pages):
        if idx >= MAX_PDF_PAGES:
            break
        try:
            parts.append(page.extract_text() or "")
        except Exception:
            continue
    text = "\n".join(parts).strip()
    if len(text) < 40:
        raise ResumeParseError(
            "PDF 没有可提取的文本，可能是扫描件。请另存为文本型 PDF 或 DOCX 后重试")
    return text


def _extract_docx(data: bytes) -> str:
    try:
        from docx import Document
    except ImportError as e:
        raise ResumeParseError(
            "缺少 python-docx 依赖，请执行 pip install -r requirements.txt") from e

    try:
        doc = Document(io.BytesIO(data))
    except Exception as e:
        raise ResumeParseError(f"DOCX 读取失败：{type(e).__name__}") from e

    parts = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _extract_plain(data: bytes) -> str:
    last_err: Exception | None = None
    for enc in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return data.decode(enc)
        except (UnicodeDecodeError, LookupError) as e:
            last_err = e
    raise ResumeParseError(
        f"文本编码无法识别：{type(last_err).__name__ if last_err else 'unknown'}")


def extract_text(filename: str, data: bytes) -> str:
    """按扩展名提取简历文本。失败抛 ResumeParseError。"""
    if not data:
        raise ResumeParseError("文件为空")
    if len(data) > MAX_UPLOAD_BYTES:
        raise ResumeParseError(
            f"文件超过 {MAX_UPLOAD_BYTES // (1024 * 1024)}MB 限制")

    suffix = PurePath(filename or "").suffix.lower()
    if suffix not in ALLOWED_EXTS:
        if suffix == ".doc":
            raise ResumeParseError(
                "旧版 .doc 不支持，请在 Word 中另存为 .docx 或 PDF")
        raise ResumeParseError(
            "仅支持 PDF / DOCX / TXT / MD 文件")

    if suffix == ".pdf":
        text = _extract_pdf(data)
    elif suffix == ".docx":
        text = _extract_docx(data)
    else:
        text = _extract_plain(data)

    text = normalize_text(text)
    if len(text) < 20:
        raise ResumeParseError("简历内容过短，无法识别")
    return text
